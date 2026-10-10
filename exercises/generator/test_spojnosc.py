"""Testy rzeczywistych plików PROSTATA-RT, nie samych funkcji wstrzykiwania."""
from __future__ import annotations
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path
import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'dane'
sys.path.insert(0,str(ROOT/'generator'))
sys.path.insert(0,str(ROOT/'walidacja'))
from generuj_dane import COLUMNS, EXTRACT, KEY_COLUMNS, text_value
from waliduj import HIDDEN_MISSING, is_missing, read_registry, read_rules, validate


def csv(name):
    return pd.read_csv(DATA/name,dtype=str,keep_default_na=False,encoding='utf-8-sig')


@pytest.fixture(scope='module')
def clean():
    return csv('rejestr_czysty.csv')


@pytest.fixture(scope='module')
def dirty():
    return read_registry(DATA/'rejestr_brudny.csv')


@pytest.fixture(scope='module')
def dictionary():
    return pd.read_csv(DATA/'slownik_danych.csv',keep_default_na=False)


def test_kontrakt_kohorty_i_slownika(clean,dictionary):
    expected='patient_id sex birth_date diagnosis_date rt_start_date weight_kg psa_dx_ng_ml gleason_primary gleason_secondary gleason_score isup_grade ct_stage cn_stage risk_group rt_technique rt_dose_gy rt_fractions adt adt_months testosterone_ng_dl psa_nadir_ng_ml psa_last_ng_ml psa_last_date bcr_phoenix gu_toxicity_max gi_toxicity_max creatinine_mg_dl hemoglobin_g_dl diabetes hypertension smoking death death_date notes'.split()
    assert COLUMNS==expected==clean.columns.tolist()==dictionary.zmienna.tolist()
    assert len(clean)==120 and clean.patient_id.tolist()==[f'PR-{i:03d}' for i in range(1,121)]
    assert set(clean.sex)=={'M'}
    assert dictionary.columns.tolist()==['zmienna','etykieta','typ','jednostka','dozwolone_wartosci','zakres_min','zakres_max','kod_braku','regula_walidacyjna','uwagi']
    assert set(dictionary.typ)<={'tekst','liczba_calkowita','liczba','data','kategoria','logiczna'}
    assert (clean.bcr_phoenix=='1').sum()==15
    assert (clean.death=='1').sum()==5
    assert clean.gu_toxicity_max.astype(int).max()==3
    assert (clean.gu_toxicity_max=='3').sum()==1
    assert clean.gi_toxicity_max.astype(int).max()<=2
    ages=[date.fromisoformat(r.rt_start_date).year-date.fromisoformat(r.birth_date).year-((date.fromisoformat(r.rt_start_date).month,date.fromisoformat(r.rt_start_date).day)<(date.fromisoformat(r.birth_date).month,date.fromisoformat(r.birth_date).day)) for r in clean.itertuples()]
    assert ages.count(89)==1 and all(55<=a<=85 or a==89 for a in ages)
    assert validate(clean,dictionary).empty


def test_klucz_odpowiada_dokladnie_faktycznym_roznicom(clean,dirty):
    key=csv('klucz_bledow.csv')
    assert key.columns.tolist()==KEY_COLUMNS
    assert len(key)==56 and len(dirty)==121
    actual={}
    for i in range(len(clean)):
        for v in clean.columns:
            if clean.at[i,v]!=dirty.at[i,v]:
                actual[(i+2,v)]=(dirty.at[i,v],clean.at[i,v])
    assert len(actual)==55
    keyed={}
    for e in key.itertuples():
        where=(int(e.wiersz_excel),e.zmienna)
        assert where not in keyed
        assert dirty.at[where[0]-2,'patient_id']==e.patient_id_w_pliku
        if where[0]<=121:
            keyed[where]=(e.wartosc_w_pliku,e.wartosc_prawdziwa)
        else:
            assert where==(122,'patient_id') and e.wartosc_w_pliku=='PR-121' and e.wartosc_prawdziwa==''
            assert e.kategoria=='duplikat'
            assert dirty.iloc[120].drop('patient_id').equals(clean.iloc[75].drop('patient_id'))
    assert actual==keyed
    assert set(key.decyzja)=={'napraw_wg_slownika','zapytanie_do_osrodka','usun_identyfikator'}


def test_rozmieszczenie_bledow_w_blokach(clean,dirty):
    # Licz rzeczywiste różnice, niezależnie od deklaracji generatora i klucza.
    per_patient=(dirty.iloc[:120]!=clean).sum(axis=1)
    counts=[int(per_patient.iloc[start:start+20].sum()) for start in range(0,120,20)]
    assert sum(counts)==55  # 56. błąd to dodatkowy rekord na końcu.
    assert all(7<=count<=12 for count in counts),counts
    summary=json.loads((DATA/'wyniki_referencyjne/podsumowanie.json').read_text())
    assert list(summary['bledy_na_bloki'].values())==counts


def test_pacjent_ma_najwyzej_dwa_bledy(clean,dirty):
    per_patient=(dirty.iloc[:120]!=clean).sum(axis=1)
    assert per_patient.max()<=2
    assert csv('klucz_bledow.csv').groupby('patient_id_w_pliku').size().max()<=2


def test_pulapki_sa_rozproszone_i_bez_bledow(clean,dirty):
    traps=csv('pulapki_nie_bledy.csv')
    positions=traps.wiersz_excel.astype(int)-2
    assert len(positions)==positions.nunique()==5
    assert positions.between(0,119).all()
    assert (positions//20).nunique()==5
    error_positions=set(csv('klucz_bledow.csv').wiersz_excel.astype(int)-2)
    assert set(positions).isdisjoint(error_positions)
    for t in traps.itertuples():
        position=int(t.wiersz_excel)-2
        assert clean.iloc[position].equals(dirty.iloc[position])
        assert clean.at[position,'patient_id']==t.patient_id


def test_wszystkie_kategorie_sa_faktycznie_w_rejestrze(dirty):
    # Lista niezależna od stałej generatora. Każdy wpis musi wskazywać komórkę pliku.
    required={'identyfikator','kodowanie','format_daty','logika_czasowa','jednostka','liczba_jako_tekst','brak_danych','zakres','spojnosc_pol','duplikat','kategoria','struktura','prywatnosc'}
    key=csv('klucz_bledow.csv')
    assert set(key.kategoria)==required
    for category in required:
        rows=key[key.kategoria==category]
        assert not rows.empty
        for e in rows.itertuples():
            assert dirty.at[int(e.wiersz_excel)-2,e.zmienna]==e.wartosc_w_pliku
    # Każdy wariant v2 występuje w pliku, nie tylko w parametryzacji walidatora.
    cases=[('patient_id','pr 014'),('sex','mężczyzna'),('sex','1'),('sex','male'),('sex','K'),
           ('psa_last_ng_ml','999'),('hemoglobin_g_dl','-9'),('testosterone_ng_dl','nd'),('creatinine_mg_dl','b.d.'),('weight_kg','brak'),('gi_toxicity_max','-'),
           ('psa_nadir_ng_ml','<0,01'),('diagnosis_date','03.04.2025'),('rt_start_date','04/03/2025'),('psa_last_date','2025.04.03'),('diagnosis_date','kwiecień 2025'),('psa_last_date','45750'),
           ('hemoglobin_g_dl','8,4 mmol/l'),('hemoglobin_g_dl','8,8 mmol/l'),('testosterone_ng_dl','0,6 nmol/l'),('weight_kg','180,7791 lb'),
           ('psa_dx_ng_ml','12,4 ng/ml'),('psa_last_ng_ml','ok. 0,2'),('rt_dose_gy','78 Gy'),('gleason_score','03.kwi'),('rt_dose_gy','60/20'),
           ('psa_dx_ng_ml','-2'),('gu_toxicity_max','7'),('rt_dose_gy','780'),('weight_kg','9'),
           ('rt_technique','vmat'),('ct_stage','T2 b'),('rt_technique','IMRT '),('smoking','tak'),('smoking','nie'),('smoking','okazjonalnie'),('notes','NT; DM2; metformina 2x1000; ASA')]
    for v,value in cases:
        indices=dirty.index[dirty[v]==value]
        assert len(indices)==1,(v,value)
        assert ((key.wiersz_excel==str(indices[0]+2)) & (key.zmienna==v) & (key.wartosc_w_pliku==value)).any()
    creat=key[(key.kategoria=='jednostka') & (key.zmienna=='creatinine_mg_dl')]
    assert len(creat)==3 and set(creat.wartosc_w_pliku)=={'88,4','106,1','70,7'}
    assert len(key[key.zmienna=='sex'])==4
    assert len(key[key.kategoria=='spojnosc_pol'])==8
    assert key[key.kategoria=='brak_danych'].zmienna.nunique()==6
    assert sum(bool(re.fullmatch(r'\d{11}',x)) for x in dirty.patient_id)==1
    assert sum(bool(re.search(r'(?<!\d)\d{11}(?!\d)',x)) for x in dirty.notes)==2
    def record(variable,category,value):
        matches=key[(key.zmienna==variable) & (key.kategoria==category) & (key.wartosc_w_pliku==value)]
        assert len(matches)==1,(variable,category,value)
        return dirty.iloc[int(matches.iloc[0].wiersz_excel)-2]
    r=record('gleason_score','spojnosc_pol','8')
    assert r.gleason_primary=='4' and r.gleason_secondary=='3' and r.gleason_score=='8'
    r=record('isup_grade','spojnosc_pol','2')
    assert r.gleason_primary=='4' and r.gleason_secondary=='3' and r.isup_grade=='2'
    r=record('risk_group','spojnosc_pol','niskie')
    assert r.psa_dx_ng_ml=='25' and r.risk_group=='niskie'
    r=record('rt_technique','spojnosc_pol','SBRT')
    assert r.rt_technique=='SBRT' and r.rt_fractions=='39'
    r=record('adt','spojnosc_pol','0')
    assert r.adt=='0' and r.adt_months=='6'
    r=record('bcr_phoenix','spojnosc_pol','0')
    assert float(r.psa_last_ng_ml)>=float(r.psa_nadir_ng_ml)+2 and r.bcr_phoenix=='0'
    r=record('death','spojnosc_pol','1')
    assert r.death=='1' and r.death_date==''
    r=record('death','spojnosc_pol','0')
    assert r.death=='0' and r.death_date!=''
    for variable,check in [('psa_last_date',lambda r:r.psa_last_date<r.rt_start_date),
                           ('death_date',lambda r:r.death_date<r.rt_start_date),
                           ('diagnosis_date',lambda r:r.diagnosis_date>r.rt_start_date)]:
        matches=key[(key.zmienna==variable) & (key.kategoria=='logika_czasowa')]
        assert len(matches)==1 and check(dirty.iloc[int(matches.iloc[0].wiersz_excel)-2])
    below=key[key.wartosc_w_pliku=='<0,01']
    assert len(below)==1 and below.iloc[0].decyzja=='zapytanie_do_osrodka'


@pytest.mark.parametrize('csv_name,xlsx_name,sheet',[
    ('slownik_danych.csv','slownik_danych.xlsx','slownik'),
    ('klucz_bledow.csv','klucz_bledow.xlsx','klucz'),
    ('rejestr_brudny.csv','rejestr_brudny.xlsx','dane')])
def test_csv_i_xlsx_sa_zgodne(csv_name,xlsx_name,sheet):
    a=read_registry(DATA/csv_name) if csv_name=='rejestr_brudny.csv' else csv(csv_name)
    b=pd.read_excel(DATA/xlsx_name,sheet_name=sheet,dtype=object,keep_default_na=False)
    assert a.columns.tolist()==b.columns.tolist()
    assert a.values.tolist()==b.map(text_value).values.tolist()
    wb=load_workbook(DATA/xlsx_name,read_only=True)
    assert wb.sheetnames==[sheet];wb.close()
    if csv_name=='rejestr_brudny.csv':
        assert (DATA/csv_name).read_bytes().startswith(b'\xef\xbb\xbf')
        assert ';' in (DATA/csv_name).read_text(encoding='utf-8-sig').splitlines()[0]


def test_pesel_w_kazdym_artefakcie_ma_bledna_kontrolna():
    found=set()
    for path in DATA.rglob('*'):
        if path.name=='manifest_sha256.csv':
            continue  # SHA256 nie jest numerem osobowym.
        texts=[]
        if path.suffix in {'.txt','.csv','.json','.yaml','.md'}:
            texts=[path.read_text(encoding='utf-8-sig')]
        elif path.suffix=='.xlsx':
            wb=load_workbook(path,read_only=True)
            texts=[' '.join(str(x or '') for row in sheet.iter_rows(values_only=True) for x in row) for sheet in wb];wb.close()
        elif path.suffix=='.zip':
            with zipfile.ZipFile(path) as z:
                texts=[z.read(n).decode('utf-8') for n in z.namelist()]
        for text in texts:
            for number in re.findall(r'(?<!\d)\d{11}(?!\d)',text):
                assert sum(int(a)*b for a,b in zip(number,[1,3,7,9,1,3,7,9,1,3,1]))%10!=0,(path,number)
                found.add(number)
    assert len(found)==63


def test_pulapki_nie_wyzwalaja_zadnego_alarmu(clean,dirty,dictionary):
    report=validate(dirty,dictionary);key=csv('klucz_bledow.csv');traps=csv('pulapki_nie_bledy.csv')
    assert len(traps)==5
    for t in traps.itertuples():
        assert dirty.at[int(t.wiersz_excel)-2,t.zmienna]==t.wartosc
        assert clean.at[int(t.wiersz_excel)-2,t.zmienna]==t.wartosc
        # Cały rekord pułapki jest wolny od alarmu i celowych błędów.
        assert report[report.patient_id==t.patient_id].empty
        assert key[key.patient_id_w_pliku==t.patient_id].empty
    records={t.zmienna:dirty.iloc[int(t.wiersz_excel)-2] for t in traps.itertuples()}
    r=records['rt_dose_gy']
    assert r.rt_technique=='SBRT' and r.rt_fractions=='5'
    assert float(r.rt_dose_gy)/5==7.25
    r=records['psa_dx_ng_ml']
    assert r.cn_stage=='N1' and r.risk_group=='wysokie'
    r=records['testosterone_ng_dl']
    assert '12 dni' in r.notes and 'pierwszym miesiącu ADT' in r.notes


def test_raport_i_pokrycie_klucza(dirty,dictionary):
    report=validate(dirty,dictionary)
    assert report.map(text_value).values.tolist()==csv('wyniki_referencyjne/walidacja_reguly_raport.csv').values.tolist()
    key=csv('klucz_bledow.csv');hits=set(zip(report.wiersz_excel.astype(int),report.zmienna))
    assert all((int(e.wiersz_excel),e.zmienna) in hits for e in key.itertuples())
    assert set(csv('wyniki_referencyjne/pokrycie_klucza.csv').wykryty)=={'1'}
    assert validate(dirty,dictionary,enabled_rules=[]).empty
    assert set(validate(dirty,dictionary,['id_format']).regula)=={'id_format'}
    catalog=read_rules()
    assert all({'id','opis','typ','zmienne'}<=r.keys() and ('wyrazenie' in r or 'zakres' in r) for r in catalog)
    assert (DATA/'reguly.yaml').read_bytes()==(ROOT/'walidacja/reguly.yaml').read_bytes()


@pytest.mark.parametrize('token',sorted(HIDDEN_MISSING|{''}))
def test_rozpoznanie_ukrytego_braku(token,clean,dictionary):
    assert is_missing(token)
    sample=clean.iloc[[0]].copy();sample.loc[:,'psa_last_ng_ml']=token
    report=validate(sample,dictionary,['brak'])
    assert ((report.zmienna=='psa_last_ng_ml') & (report.regula=='brak')).any()


@pytest.mark.parametrize('nadir,last,flag',[(.2,2.19,0),(.2,2.2,1),(.2,2.21,1),(0,2,1)])
def test_granica_phoenix(nadir,last,flag,clean,dictionary):
    sample=clean.iloc[[0]].copy()
    sample.loc[:,'psa_nadir_ng_ml']=str(nadir);sample.loc[:,'psa_last_ng_ml']=str(last);sample.loc[:,'bcr_phoenix']=str(flag)
    assert validate(sample,dictionary,['phoenix']).empty
    sample.loc[:,'bcr_phoenix']=str(1-flag)
    assert len(validate(sample,dictionary,['phoenix']))==1


@pytest.mark.parametrize('invalid',['','nd','999','-9','-1','<0,01','ok. 0,2','800'])
def test_phoenix_nie_wnioskuje_z_braku_lub_niemozliwej_wartosci(invalid,clean,dictionary):
    sample=clean.iloc[[0]].copy();sample.loc[:,'psa_last_ng_ml']=invalid
    assert validate(sample,dictionary,['phoenix']).empty


@pytest.mark.parametrize('p,s,g',[(3,3,1),(3,4,2),(4,3,3),(4,4,4),(3,5,4),(5,3,4),(4,5,5),(5,4,5),(5,5,5)])
def test_mapowanie_isup(p,s,g,clean,dictionary):
    sample=clean.iloc[[0]].copy()
    for v,value in [('gleason_primary',p),('gleason_secondary',s),('gleason_score',p+s),('isup_grade',g)]:
        sample.loc[:,v]=str(value)
    assert validate(sample,dictionary,['gleason','isup']).empty
    sample.loc[:,'isup_grade']=str(g%5+1)
    assert len(validate(sample,dictionary,['isup']))==1


@pytest.mark.parametrize('psa,grade,ct,cn,group',[
    (9.99,1,'T2a','N0','niskie'),(10,1,'T2a','N0','posrednie'),(20,1,'T2a','N0','posrednie'),(20.01,1,'T2a','N0','wysokie'),
    (5,1,'T2b','N0','posrednie'),(5,1,'T2c','N0','wysokie'),(5,1,'T1c','N1','wysokie'),(5,4,'T1c','N0','wysokie'),(12,3,'T2b','N0','posrednie')])
def test_granice_grupy_ryzyka(psa,grade,ct,cn,group,clean,dictionary):
    sample=clean.iloc[[0]].copy()
    for v,value in [('psa_dx_ng_ml',psa),('isup_grade',grade),('ct_stage',ct),('cn_stage',cn),('risk_group',group)]:
        sample.loc[:,v]=str(value)
    assert validate(sample,dictionary,['ryzyko']).empty
    sample.loc[:,'risk_group']='niskie' if group!='niskie' else 'wysokie'
    assert len(validate(sample,dictionary,['ryzyko']))==1


@pytest.mark.parametrize('dose,fractions,tech',[(78,39,'IMRT'),(76,38,'VMAT'),(60,20,'VMAT'),(62,20,'IMRT'),(70,28,'VMAT'),(36.25,5,'SBRT'),(40,5,'SBRT')])
def test_wszystkie_schematy_rt(dose,fractions,tech,clean,dictionary):
    sample=clean.iloc[[0]].copy()
    for v,value in [('rt_dose_gy',dose),('rt_fractions',fractions),('rt_technique',tech)]:
        sample.loc[:,v]=str(value)
    assert validate(sample,dictionary,['schemat_rt','sbrt']).empty
    sample.loc[:,'rt_fractions']='6'
    assert not validate(sample,dictionary,['schemat_rt','sbrt']).empty


def test_yaml_jest_rzeczywiscie_odczytywany_i_steruje_walidacja(clean,dictionary):
    with tempfile.TemporaryDirectory(prefix='.test-reguly-',dir=ROOT) as temp:
        rules=yaml.safe_load((DATA/'reguly.yaml').read_text())
        for r in rules['reguly']:
            if r['id']=='phoenix':r['przyrost']=3
            if r['id']=='zakres_weight_kg':r['zakres']['max']=60
        path=Path(temp)/'reguly.yaml';path.write_text(json.dumps(rules),encoding='utf-8')
        sample=clean.iloc[[0]].copy()
        for v,value in [('psa_nadir_ng_ml','0.2'),('psa_last_ng_ml','2.5'),('bcr_phoenix','0'),('weight_kg','80')]:sample.loc[:,v]=value
        assert not validate(sample,dictionary,['phoenix']).empty
        assert validate(sample,dictionary,['phoenix'],rules_path=path).empty
        assert not validate(sample,dictionary,['zakres_weight_kg'],rules_path=path).empty


def parse_date(text):
    for fmt in ['%Y-%m-%d','%d.%m.%Y','%d/%m/%Y','%Y.%m.%d']:
        try:return datetime.strptime(text,fmt).date().isoformat()
        except ValueError:pass
    raise AssertionError(text)


def extract_document_facts(text):
    """Niezależny parser jawnych faktów do kontroli treści; nie jest modelem NLP."""
    out={v:None for v in EXTRACT}
    out['patient_id']=re.search(r'kod rejestru: (PR-\d{3})',text)[1]
    for v,expr in [('birth_date',r'^Ur\. ([0-9./-]{10})\.'),('rt_start_date',r'(?:Pierwsza frakcja RT|RT rozpoczęto|Początek napromieniania): ([0-9./-]{10})')]:
        m=re.search(expr,text,re.M)
        if m:out[v]=parse_date(m[1])
    for v,expr in [('psa_dx_ng_ml',r'PSA przy rozpoznaniu: ([\d,]+) ng/ml'),('psa_nadir_ng_ml',r'\(nadir\): ([\d,]+) ng/ml'),('psa_last_ng_ml',r'Ostatnie PSA: ([\d,]+) ng/ml')]:
        m=re.search(expr,text)
        if m:out[v]=float(m[1].replace(',','.'))
    pair=re.search(r'(?:Gs|Gleason) (\d)\+(\d)=(\d)',text)
    if pair:
        out['gleason_primary'],out['gleason_secondary']=int(pair[1]),int(pair[2])
    if 'Sprzeczność nierozstrzygnięta' in text:
        out['isup_grade']=None
    elif 'Aneks histopatologiczny, podpisany i potwierdzony:' in text:
        out['isup_grade']=int(re.search(r'Aneks histopatologiczny, podpisany i potwierdzony: ISUP (\d)',text)[1])
    else:
        m=re.search(r'ISUP (\d)\.',text)
        if m:out['isup_grade']=int(m[1])
    for v,expr in [('ct_stage',r'Stopień kliniczny: c(T\w+)'),('cn_stage',r'Stopień kliniczny: cT\w+ c(N[01])')]:
        m=re.search(expr,text)
        if m:out[v]=m[1]
    group=re.search(r'Grupa ryzyka w karcie kwalifikacji: (niskie|pośrednie|wysokie)',text)
    if group:out['risk_group']={'niskie':'niskie','pośrednie':'posrednie','wysokie':'wysokie'}[group[1]]
    therapy=re.search(r'Technika (IMRT|VMAT|SBRT)(?:/IGRT)?; dawka na prostatę ([\d,]+) Gy w (\d+) frakcjach',text)
    if therapy:
        out['rt_technique']=therapy[1];out['rt_dose_gy']=float(therapy[2].replace(',','.'));out['rt_fractions']=int(therapy[3])
    if 'ADT zastosowano:' in text:out['adt']=1
    elif 'ADT nie stosowano.' in text:out['adt']=0
    months=re.search(r'Zaplanowany łączny czas ADT: (\d+) mies',text)
    if months:out['adt_months']=int(months[1])
    for v,label in [('gu_toxicity_max','GU'),('gi_toxicity_max','GI')]:
        m=re.search(r'CTCAE: [^\n]*\b'+label+r' G(\d)',text)
        if m:out[v]=int(m[1])
    for v,expr in [('creatinine_mg_dl',r'kreatynina ([\d,]+) (µmol/l|mg/dl)'),('hemoglobin_g_dl',r'Hb ([\d,]+) (mmol/l|g/dl)')]:
        m=re.search(expr,text)
        if m:
            val=float(m[1].replace(',','.'))
            if m[2]=='µmol/l':val/=88.4
            elif m[2]=='mmol/l':val/=.6206
            out[v]=round(val,2)
    for v,pos,neg in [('diabetes','DM t.2, leczona doustnie.','Nie stwierdzono cukrzycy.'),('hypertension','NT w wywiadzie, leczenie przewlekłe.','Nie stwierdzono nadciśnienia tętniczego.')]:
        if pos in text:out[v]=1
        elif neg in text:out[v]=0
    for literal,value in [('Nigdy nie palił.','never'),('Nie pali od 10 lat.','former'),('Nadal pali papierosy.','current')]:
        if literal in text:out['smoking']=value
    if out['psa_nadir_ng_ml'] is not None and out['psa_last_ng_ml'] is not None:
        out['bcr_phoenix']=int(out['psa_last_ng_ml']>=out['psa_nadir_ng_ml']+2)
    return out


@pytest.mark.parametrize('index',range(1,21))
def test_dokument_z_kluczem_i_dowodami(index,clean):
    document=f'DOK-{index:02d}.txt';text=(DATA/'dokumenty'/document).read_text(encoding='utf-8')
    assert 150<=len(text.split())<=400
    assert 'OŚRODEK FIKCYJNY' in text and f'Fikcyjny-{index:02d}' in text
    gold=csv('ekstrakcja_klucz.csv').set_index('dokument').loc[document]
    assert gold.patient_id==f'PR-{100+index:03d}'
    source=clean.set_index('patient_id',drop=False).loc[gold.patient_id]
    facts=extract_document_facts(text)
    for v in EXTRACT:
        if gold[v]=='':
            assert facts[v] is None,(document,v,facts[v])
        else:
            assert gold[v]==source[v],(document,v,gold[v],source[v])
            if isinstance(facts[v],(int,float)):
                assert facts[v]==pytest.approx(float(gold[v]),abs=.005),(document,v)
            else:
                assert facts[v]==gold[v],(document,v)
    evidence=csv('ekstrakcja_dowody.csv');evidence=evidence[evidence.dokument==document]
    assert set(evidence.zmienna)=={v for v in EXTRACT if gold[v]!=''}
    for e in evidence.itertuples():
        assert e.fragment in text and e.wartosc_klucza==gold[e.zmienna]


@pytest.mark.parametrize('index',range(21,61))
def test_nowe_dokumenty_z_kluczem_i_dowodami(index,clean):
    document=f'DOK-{index:02d}.txt';text=(DATA/'dokumenty'/document).read_text(encoding='utf-8')
    assert 150<=len(text.split())<=400
    assert 'OŚRODEK FIKCYJNY' in text and f'Fikcyjny-{index:02d}' in text
    gold=csv('ekstrakcja_klucz.csv').set_index('dokument').loc[document]
    assert gold.patient_id==f'PR-{40+index:03d}'
    source=clean.set_index('patient_id',drop=False).loc[gold.patient_id]
    facts=extract_document_facts(text)
    for v in EXTRACT:
        if gold[v]=='':
            assert facts[v] is None,(document,v,facts[v])
        else:
            assert gold[v]==source[v],(document,v,gold[v],source[v])
            if isinstance(facts[v],(int,float)):
                assert facts[v]==pytest.approx(float(gold[v]),abs=.005),(document,v)
            else:
                assert facts[v]==gold[v],(document,v)
    evidence=csv('ekstrakcja_dowody.csv');evidence=evidence[evidence.dokument==document]
    assert set(evidence.zmienna)=={v for v in EXTRACT if gold[v]!=''}
    for e in evidence.itertuples():
        assert e.fragment in text and e.wartosc_klucza==gold[e.zmienna]


def test_sha256_nienaruszone_sprzed_a4():
    sha_file=ROOT.parent/'plan'/'sha256_przed_A4.txt'
    assert sha_file.is_file(),f"Brak pliku sum kontrolnych: {sha_file}"
    lines=[line.strip() for line in sha_file.read_text(encoding='utf-8').splitlines() if line.strip()]
    base=ROOT.parent
    for line in lines:
        expected_sha,rel_path=line.split()
        target=base/rel_path
        assert target.is_file(),f"Brak pliku bazowego: {rel_path}"
        actual_sha=hashlib.sha256(target.read_bytes()).hexdigest()
        assert actual_sha==expected_sha,(
            f"Naruszenie sumy kontrolnej sprzed A4 dla {rel_path}: oczekiwano {expected_sha}, otrzymano {actual_sha}"
        )


def test_schemat_braki_i_rozstrzyganie_sprzecznosci():
    expected='patient_id birth_date rt_start_date psa_dx_ng_ml gleason_primary gleason_secondary isup_grade ct_stage cn_stage risk_group rt_technique rt_dose_gy rt_fractions adt adt_months psa_nadir_ng_ml psa_last_ng_ml bcr_phoenix gu_toxicity_max gi_toxicity_max creatinine_mg_dl hemoglobin_g_dl diabetes hypertension smoking'.split()
    assert EXTRACT==expected
    schema=json.loads((DATA/'ekstrakcja_schemat.json').read_text())
    assert schema['required']==expected and list(schema['properties'])==expected and not schema['additionalProperties']
    gold=csv('ekstrakcja_klucz.csv')
    assert gold.columns.tolist()==['dokument']+expected and len(gold)==60
    for row in gold.to_dict('records'):
        for v in EXTRACT:
            value=row[v];spec=schema['properties'][v]
            assert 'null' in spec['type']
            if value=='':continue
            if 'enum' in spec:assert (int(value) if 'integer' in spec['type'] else value) in spec['enum']
            if 'pattern' in spec:assert re.fullmatch(spec['pattern'],value)
            if 'minimum' in spec:assert spec['minimum']<=float(value)<=spec['maximum']
    d=csv('ekstrakcja_trudnosci.csv')
    assert set(d.typ_trudnosci)>={'data','jednostka','negacja','sprzecznosc','brak','oznaczalnosc'}
    for r in gold.to_dict('records'):
        for v in EXTRACT:
            if r[v]=='':assert ((d.dokument==r['dokument']) & (d.zmienna==v)).any()
    g=gold.set_index('dokument')
    assert g.loc['DOK-16.txt','isup_grade']=='3'
    assert g.loc['DOK-17.txt','isup_grade']==''
    assert g.loc['DOK-18.txt','psa_last_ng_ml']==g.loc['DOK-18.txt','bcr_phoenix']==''
    assert g.loc['DOK-19.txt','birth_date']==''
    assert g.loc['DOK-14.txt','birth_date']!=''


def test_zip_karta_manifest_i_scenariusz():
    with zipfile.ZipFile(DATA/'dokumenty.zip') as z:
        assert z.namelist()==[f'DOK-{i:02d}.txt' for i in range(1,61)]
        for name in z.namelist():assert z.read(name)==(DATA/'dokumenty'/name).read_bytes()
    wb=load_workbook(DATA/'karta_znalezisk.xlsx',read_only=True)
    assert list(next(wb.active.iter_rows(values_only=True)))==['patient_id','zmienna','problem','decyzja'];wb.close()
    manifest=csv('manifest_sha256.csv')
    paths={p.relative_to(DATA).as_posix() for p in DATA.rglob('*') if p.is_file() and p.name!='manifest_sha256.csv'}
    assert set(manifest.plik)==paths
    for item in manifest.itertuples():
        data=(DATA/item.plik).read_bytes()
        assert hashlib.sha256(data).hexdigest()==item.sha256 and len(data)==int(item.bajty)
    s=json.loads((DATA/'scenariusz.json').read_text())
    assert s['nazwa']=='PROSTATA-RT' and s['syntetyczne'] and len(s['definicje'])>=6
    assert not list(DATA.glob('*kardioonko*'))
    if (ROOT/'_archiwum_kardioonko').exists():
        assert (ROOT/'_archiwum_kardioonko/dane/rejestr_kardioonko_czysty.csv').is_file()


def test_odtworzenie_od_zera_i_deterministycznosc():
    with tempfile.TemporaryDirectory(prefix='.test-generator-',dir=ROOT) as temp:
        target=Path(temp)
        for folder,name in [('generator','generuj_dane.py'),('walidacja','waliduj.py')]:
            (target/folder).mkdir();shutil.copyfile(ROOT/folder/name,target/folder/name)
        command=[sys.executable,'-B',str(target/'generator/generuj_dane.py')]
        env={**os.environ,'PYTHONDONTWRITEBYTECODE':'1'}
        result=subprocess.run(command,capture_output=True,text=True,env=env,timeout=60)
        assert result.returncode==0,result.stderr
        assert json.loads(result.stdout)['bledy']==56
        def files(base):return {p.relative_to(base).as_posix():p.read_bytes() for p in base.rglob('*') if p.is_file()}
        first=files(target/'dane')
        assert first==files(DATA)
        result=subprocess.run(command,capture_output=True,text=True,env=env,timeout=60)
        assert result.returncode==0,result.stderr
        assert files(target/'dane')==first
        assert (target/'README_dane.md').read_bytes()==(ROOT/'README_dane.md').read_bytes()


def test_walidator_cli_csv_i_xlsx():
    with tempfile.TemporaryDirectory(prefix='.test-walidacja-',dir=ROOT) as temp:
        outputs=[]
        for extension in ['csv','xlsx']:
            output=Path(temp)/f'{extension}.csv'
            result=subprocess.run([sys.executable,'-B',str(ROOT/'walidacja/waliduj.py'),str(DATA/f'rejestr_brudny.{extension}'),'--raport',str(output)],capture_output=True,text=True,env={**os.environ,'PYTHONDONTWRITEBYTECODE':'1'},timeout=60)
            assert result.returncode==0,result.stderr
            outputs.append(output.read_bytes())
        assert outputs[0]==outputs[1]==(DATA/'wyniki_referencyjne/walidacja_reguly_raport.csv').read_bytes()
